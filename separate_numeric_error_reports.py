from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent
ERROR_DESCRIPTION_COLUMN = "ERROR_DESCRIPTION"
LEGACY_DESTINATION_FOLDER = "numeric_column_errors"
DESTINATION_FOLDER = "other_errors"


def _restore_legacy_reports(errors_dir: Path) -> None:
    legacy_destination_dir = errors_dir / LEGACY_DESTINATION_FOLDER
    for legacy_report_path in sorted(legacy_destination_dir.glob("*.xlsx")):
        if legacy_report_path.name.startswith("~$"):
            continue
        restored_path = errors_dir / legacy_report_path.name
        if restored_path.exists():
            print(f"SKIPPED {legacy_report_path.name}: file already exists in {errors_dir}")
            continue
        try:
            shutil.move(str(legacy_report_path), str(restored_path))
        except PermissionError:
            print(f"SKIPPED {legacy_report_path.name}: file is in use")


def _has_other_description(
    report_path: Path,
    numeric_error_pattern: re.Pattern[str],
) -> bool | None:
    try:
        workbook = load_workbook(report_path, read_only=True, data_only=True)
    except Exception as exc:
        print(f"SKIPPED {report_path.name}: could not read workbook ({exc})")
        return None

    try:
        sheet = workbook.active
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        try:
            description_index = headers.index(ERROR_DESCRIPTION_COLUMN) + 1
        except ValueError:
            return None
        return any(
            not numeric_error_pattern.search(str(row[0]).strip())
            for row in sheet.iter_rows(
                min_row=2,
                min_col=description_index,
                max_col=description_index,
                values_only=True,
            )
            if row[0] is not None and str(row[0]).strip()
        )
    finally:
        workbook.close()


def _move_other_reports(
    errors_dir: Path,
    destination_dir: Path,
    numeric_error_pattern: re.Pattern[str],
) -> int:
    moved_count = 0
    for report_path in sorted(errors_dir.glob("*.xlsx")):
        if report_path.name.startswith("~$"):
            continue
        has_other_error = _has_other_description(report_path, numeric_error_pattern)
        if has_other_error is not True:
            continue

        destination_path = destination_dir / report_path.name
        if destination_path.exists():
            print(f"SKIPPED {report_path.name}: destination already exists")
            continue
        try:
            shutil.move(str(report_path), str(destination_path))
        except PermissionError:
            print(f"SKIPPED {report_path.name}: file is in use")
            continue
        moved_count += 1
        print(f"MOVED {report_path.name} -> {destination_dir.name}/")
    return moved_count


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Move reports with errors beyond numeric conversion errors into a separate folder."
    )
    parser.add_argument(
        "--errors-dir",
        type=Path,
        default=PROJECT_DIR / "load_errors",
        help="folder containing error report workbooks",
    )
    args = parser.parse_args()

    errors_dir = args.errors_dir.resolve()
    destination_dir = errors_dir / DESTINATION_FOLDER
    destination_dir.mkdir(parents=True, exist_ok=True)
    numeric_error_pattern = re.compile(
        r"Numeric column '[^']+' contains non-numeric value '[^']*'",
        re.IGNORECASE,
    )

    _restore_legacy_reports(errors_dir)
    moved_count = _move_other_reports(errors_dir, destination_dir, numeric_error_pattern)
    print(f"Moved {moved_count} report(s) into {destination_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())