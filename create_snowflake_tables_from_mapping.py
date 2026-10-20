from __future__ import annotations

import argparse
import os
import re
import sys
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from fourd_to_snowflake import SnowflakeTargetConfig, quote_snowflake_identifier


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_DATABASE = "DW_AGAM_DEV"
DEFAULT_SCHEMA = "UW_RAW_STG"
REQUIRED_HEADERS = {"target column", "snowflake datatype"}
DATATYPE_PATTERN = re.compile(
    r"(?:NUMBER|NUMERIC|DECIMAL|INT|INTEGER|BIGINT|SMALLINT|TINYINT|BYTEINT|"
    r"FLOAT|FLOAT4|FLOAT8|DOUBLE(?: PRECISION)?|REAL|VARCHAR|CHAR|CHARACTER|"
    r"STRING|TEXT|BINARY|BOOLEAN|BOOL|DATE|TIME|TIMESTAMP(?:_NTZ|_LTZ|_TZ)?|"
    r"VARIANT|OBJECT|ARRAY)(?:\s*\(\s*\d+(?:\s*,\s*\d+)?\s*\))?",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ColumnDefinition:
    name: str
    datatype: str
    is_primary_key: bool


@dataclass(frozen=True)
class TableDefinition:
    name: str
    columns: tuple[ColumnDefinition, ...]
    worksheet: str


def _normalized_header(value: Any) -> str:
    return str(value).strip().casefold() if value is not None else ""


def _find_header_row(worksheet: Any) -> tuple[int, dict[str, int]] | None:
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 30), values_only=True),
        start=1,
    ):
        headers = {_normalized_header(value): index for index, value in enumerate(row)}
        if REQUIRED_HEADERS.issubset(headers):
            return row_number, headers
    return None


def _row_value(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers.get(name)
    return row[index] if index is not None and index < len(row) else None


def _parse_mapping_row(
    worksheet_name: str,
    row_number: int,
    row: tuple[Any, ...],
    headers: dict[str, int],
) -> tuple[ColumnDefinition | None, str | None]:
    raw_column = _row_value(row, headers, "target column")
    if raw_column is None or not str(raw_column).strip():
        return None, None

    raw_datatype = _row_value(row, headers, "snowflake datatype")
    if raw_datatype is None or not str(raw_datatype).strip():
        raise ValueError(f"{worksheet_name!r}, row {row_number}: missing Snowflake Datatype")

    column_name = str(raw_column).strip()
    datatype = re.sub(r"\s+", " ", str(raw_datatype).strip())
    if not DATATYPE_PATTERN.fullmatch(datatype):
        raise ValueError(
            f"{worksheet_name!r}, row {row_number}: unsupported or unsafe "
            f"Snowflake Datatype {datatype!r}"
        )

    raw_table = _row_value(row, headers, "target table")
    table_name = str(raw_table).strip() if raw_table is not None else None
    raw_constraints = _row_value(row, headers, "constraints")
    constraints = str(raw_constraints).strip() if raw_constraints is not None else ""
    return (
        ColumnDefinition(
            name=column_name,
            datatype=datatype,
            is_primary_key=bool(re.search(r"\bPK\b", constraints, re.IGNORECASE)),
        ),
        table_name or None,
    )


def _read_worksheet_definition(worksheet: Any) -> TableDefinition | None:
    header_info = _find_header_row(worksheet)
    if header_info is None:
        print(f"SKIPPED worksheet {worksheet.title!r}: mapping headers not found")
        return None

    header_row, headers = header_info
    table_names: set[str] = set()
    columns: list[ColumnDefinition] = []
    seen_columns: set[str] = set()
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        column, table_name = _parse_mapping_row(
            worksheet.title, row_number, row, headers
        )
        if column is None:
            continue
        normalized_column = column.name.casefold()
        if normalized_column in seen_columns:
            raise ValueError(
                f"{worksheet.title!r}: duplicate Target Column {column.name!r}"
            )
        seen_columns.add(normalized_column)
        columns.append(column)
        if table_name:
            table_names.add(table_name)

    if not columns:
        print(f"SKIPPED worksheet {worksheet.title!r}: no target columns found")
        return None
    if len(table_names) > 1:
        raise ValueError(
            f"{worksheet.title!r}: multiple Target Table values: "
            + ", ".join(sorted(table_names))
        )

    table_name = next(iter(table_names), worksheet.title)
    return TableDefinition(table_name, tuple(columns), worksheet.title)


def read_table_definitions(workbook_path: Path) -> list[TableDefinition]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Mapping workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    definitions: list[TableDefinition] = []
    try:
        for worksheet in workbook.worksheets:
            definition = _read_worksheet_definition(worksheet)
            if definition is not None:
                definitions.append(definition)
    finally:
        workbook.close()

    if not definitions:
        raise ValueError("No worksheet contained usable table and column definitions")
    return definitions


def build_create_table_sql(
    definition: TableDefinition,
    database: str,
    schema: str,
    replace_existing: bool = False,
) -> str:
    qualified_table = ".".join(
        quote_snowflake_identifier(identifier)
        for identifier in (database, schema, definition.name)
    )
    columns = [
        f"  {quote_snowflake_identifier(column.name)} {column.datatype}"
        for column in definition.columns
    ]
    primary_keys = [column.name for column in definition.columns if column.is_primary_key]
    if primary_keys:
        quoted_keys = ", ".join(quote_snowflake_identifier(name) for name in primary_keys)
        columns.append(f"  PRIMARY KEY ({quoted_keys})")

    create = "CREATE OR REPLACE TABLE" if replace_existing else "CREATE TABLE"
    return f"{create} {qualified_table} (\n" + ",\n".join(columns) + "\n)"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Create one empty Snowflake table per worksheet in a mapping workbook."
    )
    parser.add_argument("workbook", type=Path, help="Excel mapping workbook (.xlsx or .xlsm)")
    parser.add_argument("--database", default=DEFAULT_DATABASE)
    parser.add_argument("--schema", default=DEFAULT_SCHEMA)
    parser.add_argument(
        "--replace",
        action="store_true",
        help="replace existing tables (drops their current data)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="print the CREATE TABLE statements without connecting to Snowflake",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.workbook.suffix.casefold() not in {".xlsx", ".xlsm"}:
        print("Input must be an .xlsx or .xlsm workbook", file=sys.stderr)
        return 2

    try:
        definitions = read_table_definitions(args.workbook.resolve())
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    statements = [
        build_create_table_sql(definition, args.database, args.schema, args.replace)
        for definition in definitions
    ]
    if args.dry_run:
        for definition, statement in zip(definitions, statements):
            print(f"-- worksheet: {definition.worksheet}; table: {definition.name}")
            print(statement + ";")
        print(f"Dry run: {len(statements)} table(s) would be created.")
        return 0

    os.environ.setdefault("SNOWFLAKE_DATABASE", args.database)
    os.environ.setdefault("SNOWFLAKE_SCHEMA", args.schema)
    try:
        config = replace(
            SnowflakeTargetConfig.from_env(),
            database=args.database,
            schema=args.schema,
        )
        connection = config.connect()
    except (RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    failed_count = 0
    cursor = connection.cursor()
    try:
        database_name = quote_snowflake_identifier(args.database)
        schema_name = quote_snowflake_identifier(args.schema)
        cursor.execute(f"CREATE SCHEMA IF NOT EXISTS {database_name}.{schema_name}")
        for definition, statement in zip(definitions, statements):
            try:
                cursor.execute(statement)
                print(f"CREATED {args.database}.{args.schema}.{definition.name}")
            except Exception as exc:
                failed_count += 1
                print(
                    f"FAILED {definition.worksheet!r} -> {definition.name!r}: {exc}",
                    file=sys.stderr,
                )
    finally:
        cursor.close()
        connection.close()

    created_count = len(definitions) - failed_count
    print(f"Finished: {created_count} created, {failed_count} failed.")
    return 1 if failed_count else 0


if __name__ == "__main__":
    raise SystemExit(main())